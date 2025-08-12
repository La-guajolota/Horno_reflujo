import asyncio
import websockets
import json
import openpyxl
from openpyxl.styles import Font, PatternFill
from datetime import datetime
from colorama import Fore, Style, init

# Inicializa colorama para que los colores funcionen en todas las terminales
init(autoreset=True)

# Configuración del dispositivo ESP8266 (solo un horno)
DEVICE = {"ip": "192.168.1.150", "name": "Horno de Reflujo"}
ESP_PORT = 80
WEBSOCKET_PATH = "/ws"

# Nombre del archivo Excel único
EXCEL_FILENAME = "regresion_horno.xlsx"

# Variables globales para el archivo Excel
workbook = None
current_sheet = None

def setup_excel():
    global workbook, current_sheet
    
    try:
        # Intentar cargar el archivo existente
        workbook = openpyxl.load_workbook(EXCEL_FILENAME)
        print(f"{Fore.YELLOW}⚠️ Archivo Excel '{EXCEL_FILENAME}' ya existe. Creando nueva hoja.{Style.RESET_ALL}")
    except FileNotFoundError:
        # Crear nuevo archivo si no existe
        workbook = openpyxl.Workbook()
        # Eliminar la hoja por defecto creada automáticamente
        if 'Sheet' in workbook.sheetnames:
            del workbook['Sheet']
        print(f"{Fore.GREEN}✅ Nuevo archivo Excel '{EXCEL_FILENAME}' creado.{Style.RESET_ALL}")
    
    # Crear nueva hoja para esta sesión con nombre basado en fecha/hora
    session_name = datetime.now().strftime("%Y-%m-%d_%H-%M-%S")
    current_sheet = workbook.create_sheet(title=session_name)
    
    # Cabeceras
    current_sheet['A1'] = "Timestamp"
    current_sheet['B1'] = "Temperatura (°C)"
    current_sheet['C1'] = "Estado"
    
    # Estilos para cabeceras
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="0070C0", end_color="0070C0", fill_type="solid")
    
    for col in ['A', 'B', 'C']:
        current_sheet[f'{col}1'].font = header_font
        current_sheet[f'{col}1'].fill = header_fill
        
    # Guardar el archivo
    workbook.save(EXCEL_FILENAME)
    print(f"{Fore.GREEN}✅ Nueva hoja '{session_name}' creada para esta sesión.{Style.RESET_ALL}")
        
    return current_sheet

async def log_to_excel(timestamp, temp, estado):
    """Registra datos en la hoja actual de Excel"""
    global current_sheet, workbook
    
    if current_sheet is None:
        return
        
    try:
        row_data = [
            timestamp,
            float(temp),
            estado
        ]
        current_sheet.append(row_data)
        workbook.save(EXCEL_FILENAME)
    except Exception as e:
        print(f"{Fore.RED}❌ Error al escribir en Excel: {e}{Style.RESET_ALL}")

async def connect_to_esp():
    """Establece conexión WebSocket con el dispositivo ESP8266"""
    ip = DEVICE['ip']
    name = DEVICE['name']
    uri = f"ws://{ip}:{ESP_PORT}{WEBSOCKET_PATH}"
    device_tag = f"{Fore.LIGHTBLACK_EX}[{name}]{Style.RESET_ALL}"
    
    print(f"{device_tag}{Fore.CYAN} 🚀 Intentando conectar a {ip}{Style.RESET_ALL}")

    while True:
        try:
            async with websockets.connect(uri) as websocket:
                print(f"{device_tag}{Fore.GREEN} ✅ Conexión establecida. Esperando datos...{Style.RESET_ALL}")
                
                while True:
                    try:
                        message = await websocket.recv()
                        
                        try:
                            data = json.loads(message)
                            temp = data.get("temp", "N/A")
                            estado = data.get("estado", "N/A")

                            current_time = datetime.now()
                            timestamp_str = current_time.strftime("%Y-%m-%d %H:%M:%S")

                            # Selección de emojis y colores
                            status_emoji = "⏳"
                            temp_color = Fore.WHITE
                            state_color = Fore.WHITE

                            if "IDLE" in estado:
                                status_emoji = "😴"
                                state_color = Fore.BLUE
                            elif "PREHEAT" in estado:
                                status_emoji = "🔥"
                                state_color = Fore.YELLOW
                            elif "SOAK" in estado:
                                status_emoji = "♨️"
                                state_color = Fore.MAGENTA
                            elif "HEATUP" in estado:
                                status_emoji = "⬆️"
                                state_color = Fore.LIGHTRED_EX
                            elif "REFLOW" in estado:
                                status_emoji = "✨"
                                state_color = Fore.RED
                            elif "COOLDOWN" in estado:
                                status_emoji = "❄️"
                                state_color = Fore.CYAN
                            
                            try:
                                temp_val = float(temp)
                                if temp_val < 50:
                                    temp_color = Fore.BLUE
                                elif temp_val < 150:
                                    temp_color = Fore.GREEN
                                elif temp_val < 200:
                                    temp_color = Fore.YELLOW
                                else:
                                    temp_color = Fore.RED
                            except ValueError:
                                pass

                            # Mostrar datos con prefijo de dispositivo
                            print(f"{device_tag} {timestamp_str} | {status_emoji} Estado: {state_color}{estado}{Fore.RESET} | {temp_color}🌡️ {temp}°C{Style.RESET_ALL}")
                            
                            # Registrar en Excel
                            await log_to_excel(timestamp_str, temp, estado)

                        except json.JSONDecodeError:
                            print(f"{device_tag}{Fore.YELLOW} ⚠️ Mensaje no JSON: {message}{Style.RESET_ALL}")

                    except websockets.exceptions.ConnectionClosed:
                        print(f"{device_tag}{Fore.RED} ❌ Conexión cerrada. Reconectando en 5s...{Style.RESET_ALL}")
                        break
                    except Exception as e:
                        print(f"{device_tag}{Fore.RED} ❌ Error: {e}. Reconectando en 5s...{Style.RESET_ALL}")
                        break
            
        except (ConnectionRefusedError, asyncio.TimeoutError):
            print(f"{device_tag}{Fore.RED} ❌ Error de conexión. Reintentando en 5s...{Style.RESET_ALL}")
        except Exception as e:
            print(f"{device_tag}{Fore.RED} ❌ Error inesperado: {e}. Reconectando en 5s...{Style.RESET_ALL}")
        
        await asyncio.sleep(5)

async def main():
    """Función principal para manejar el dispositivo"""
    # Configurar Excel con nueva hoja
    setup_excel()
    
    print(f"\n{Fore.YELLOW}💡 Presiona Ctrl+C para salir.{Style.RESET_ALL}")
    
    try:
        await connect_to_esp()
    except KeyboardInterrupt:
        print(f"\n{Fore.CYAN}👋 Desconectando dispositivo...{Style.RESET_ALL}")
    finally:
        # Cerrar el archivo Excel
        if workbook:
            try:
                workbook.save(EXCEL_FILENAME)
                workbook.close()
                print(f"{Fore.GREEN}✅ Archivo {EXCEL_FILENAME} guardado y cerrado correctamente{Style.RESET_ALL}")
            except Exception as e:
                print(f"{Fore.RED}❌ Error cerrando Excel: {e}{Style.RESET_ALL}")

if __name__ == "__main__":
    asyncio.run(main())
